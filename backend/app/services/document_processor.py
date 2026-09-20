"""
Document Processor for EduPath.

Extracts text and metadata from uploaded files (PDF, DOCX, TXT, MD, JSON)
and raw portfolio descriptions using pypdf and python-docx with graceful fallback.
"""

from abc import ABC, abstractmethod
import io
import json
from typing import Any, Dict, Optional

from app.utils.logger import logger

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None


class BaseDocumentProcessor(ABC):
    """Abstract interface for extracting text and metadata from uploaded documents."""

    @abstractmethod
    async def extract_text(self, file_content: bytes, filename: str) -> str:
        """Extract raw text content from uploaded document bytes."""
        pass

    @abstractmethod
    async def extract_metadata(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Extract file metadata (e.g. page count, author, size)."""
        pass


class RealDocumentProcessor(BaseDocumentProcessor):
    """Real document processor using pypdf, python-docx, and UTF-8 text decoding."""

    async def extract_text(self, file_content: bytes, filename: str) -> str:
        ext = filename.lower().split(".")[-1] if "." in filename else ""

        if ext == "pdf" and PdfReader is not None:
            try:
                reader = PdfReader(io.BytesIO(file_content))
                extracted_pages = []
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        extracted_pages.append(text)
                full_text = "\n".join(extracted_pages)
                if full_text.strip():
                    return full_text
            except Exception as e:
                logger.warning(f"pypdf extraction failed for {filename}: {e}")

        elif ext in ("docx", "doc") and docx is not None:
            try:
                doc = docx.Document(io.BytesIO(file_content))
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                full_text = "\n".join(paragraphs)
                if full_text.strip():
                    return full_text
            except Exception as e:
                logger.warning(f"python-docx extraction failed for {filename}: {e}")

        # Fallback to plain text decoding (TXT, MD, JSON, or unparsed text)
        try:
            decoded = file_content.decode("utf-8", errors="replace")
            if ext == "json":
                try:
                    data = json.loads(decoded)
                    return json.dumps(data, indent=2)
                except Exception:
                    pass
            return decoded
        except Exception as e:
            logger.error(f"Failed to decode text for {filename}: {e}")
            return f"Raw file content from {filename}"

    async def extract_metadata(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        ext = filename.lower().split(".")[-1] if "." in filename else ""
        meta = {
            "filename": filename,
            "extension": ext,
            "size_bytes": len(file_content),
        }

        if ext == "pdf" and PdfReader is not None:
            try:
                reader = PdfReader(io.BytesIO(file_content))
                meta["page_count"] = len(reader.pages)
            except Exception:
                meta["page_count"] = 1
        elif ext in ("docx", "doc") and docx is not None:
            try:
                doc = docx.Document(io.BytesIO(file_content))
                meta["paragraph_count"] = len(doc.paragraphs)
            except Exception:
                meta["paragraph_count"] = 1

        return meta


def get_document_processor() -> BaseDocumentProcessor:
    """Returns an active instance of RealDocumentProcessor."""
    return RealDocumentProcessor()
